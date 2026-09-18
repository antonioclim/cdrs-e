#!/usr/bin/env python3
from __future__ import annotations

from pathlib import Path
import subprocess

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
REPORTS = ROOT / "reports"
REFERENCE = REPORTS / "software_report_reference.docx"
SOURCE = REPORTS / "CDRS_E_PHASE_XLII_COMPANION_SOFTWARE_REPORT.md"
OUTPUT = REPORTS / "CDRS_E_PHASE_XLII_COMPANION_SOFTWARE_REPORT.docx"


def set_font(style, name: str, size: float, bold: bool = False) -> None:
    style.font.name = name
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)


def page_field(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Antonio Clim • Phase XLII • 12 September 2026 • ")
    run.font.name = "Carlito"
    run.font.size = Pt(8)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


def build_reference() -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.0)
    sec.bottom_margin = Cm(2.0)
    sec.left_margin = Cm(2.2)
    sec.right_margin = Cm(2.2)
    sec.header_distance = Cm(0.8)
    sec.footer_distance = Cm(0.8)
    header = sec.header.paragraphs[0]
    header.text = "CDRS-E  |  Phase XLII, companion software productisation"
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in header.runs:
        run.font.name = "Carlito"
        run.font.size = Pt(8)
        run.bold = True
    page_field(sec.footer.paragraphs[0])
    styles = doc.styles
    set_font(styles["Normal"], "Carlito", 10.2)
    styles["Normal"].paragraph_format.line_spacing = 1.08
    styles["Normal"].paragraph_format.space_after = Pt(4)
    styles["Normal"].paragraph_format.widow_control = True
    for name, size, before, after in [
        ("Title", 18, 0, 12),
        ("Subtitle", 11, 0, 8),
        ("Heading 1", 14, 12, 5),
        ("Heading 2", 11.5, 9, 4),
        ("Heading 3", 10.5, 7, 3),
    ]:
        set_font(styles[name], "Carlito", size, bold=name != "Subtitle")
        styles[name].paragraph_format.space_before = Pt(before)
        styles[name].paragraph_format.space_after = Pt(after)
        styles[name].paragraph_format.keep_with_next = True
    set_font(styles["Caption"], "Carlito", 8.8)
    styles["Caption"].font.italic = True
    styles["Caption"].paragraph_format.space_after = Pt(6)
    doc.core_properties.author = "Antonio Clim"
    doc.core_properties.title = "CDRS-E Companion Software Productisation"
    doc.core_properties.subject = "Phase XLII software architecture, verification and release gates"
    doc.core_properties.keywords = "CDRS-E; research software; reproducibility; certificates; multivariate time series"
    doc.save(REFERENCE)


def main() -> int:
    build_reference()
    subprocess.run(
        [
            "pandoc",
            SOURCE.name,
            "--from=markdown+tex_math_dollars+tex_math_single_backslash+pipe_tables+yaml_metadata_block",
            "--to=docx",
            "--standalone",
            "--reference-doc",
            str(REFERENCE),
            "--resource-path",
            str(REPORTS),
            "--output",
            str(OUTPUT),
        ],
        cwd=REPORTS,
        check=True,
    )
    print(OUTPUT)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
